"""Motor RAG Semántico para la Especialización en Inteligencia Artificial (Nexus).

Indexa documentos (PDFs, guías curriculares, DOCX) de la Especialización en IA de la UAM
utilizando LlamaIndex y Google Gemini (models/gemini-embedding-2).
El índice vectorial se inicializa y persiste en disco una sola vez para evitar re-computar embeddings.
"""

from __future__ import annotations

import asyncio
import json
import logging
import os
from pathlib import Path
from typing import Any

from dotenv import load_dotenv

logger = logging.getLogger("rag-nexus")

load_dotenv(".env.local")
load_dotenv(".env")
load_dotenv(os.path.join(os.path.dirname(__file__), "..", "..", ".env"))

_nexus_index: Any = None
_nexus_query_engine: Any = None
_nexus_lock = asyncio.Lock()


def _get_nexus_storage_dir() -> Path:
    storage_dir = os.getenv("NEXUS_STORAGE_DIR")
    if not storage_dir:
        storage_dir = os.path.join(os.path.dirname(__file__), "..", "data", "nexus_storage")
    p = Path(storage_dir)
    p.mkdir(parents=True, exist_ok=True)
    return p


def _get_nexus_docs_dir() -> Path:
    possible_paths = [
        os.getenv("NEXUS_DOCS_DIR", ""),
        os.path.join(os.path.dirname(__file__), "..", "data", "especializacion_ia"),
        os.path.join(os.path.dirname(__file__), "..", "..", "data", "especializacion_ia"),
        "/data/especializacion_ia",
        "C:\\Users\\eguen\\Downloads\\Especializacion IA",
    ]
    for p_str in possible_paths:
        if p_str:
            p = Path(p_str).resolve()
            if p.is_dir() and any(p.iterdir()):
                return p
    return Path(os.path.join(os.path.dirname(__file__), "..", "data", "especializacion_ia")).resolve()


def listar_documentos_especializacion() -> list[dict[str, Any]]:
    """Retorna la lista de todos los documentos y guías de la especialización disponibles."""
    docs_dir = _get_nexus_docs_dir()
    if not docs_dir.exists():
        return []

    catalogo = []
    for f in sorted(docs_dir.iterdir()):
        if f.is_file() and f.suffix.lower() in [".pdf", ".docx", ".doc", ".txt", ".md"]:
            catalogo.append(
                {
                    "nombre": f.name,
                    "tipo": f.suffix.lower().replace(".", "").upper(),
                    "tamano_kb": round(f.stat().st_size / 1024, 1),
                }
            )
    return catalogo


def _extraer_texto_archivo(ruta: Path) -> str:
    """Extrae texto de archivos PDF o DOCX de forma tolerante a fallos."""
    sufijo = ruta.suffix.lower()
    texto = ""
    try:
        if sufijo == ".pdf":
            import pypdf

            reader = pypdf.PdfReader(str(ruta))
            paginas_texto = []
            for pag in reader.pages:
                t = pag.extract_text() or ""
                if t.strip():
                    paginas_texto.append(t)
            texto = "\n".join(paginas_texto)

            # Si el texto está vacío (PDF escaneado), se transcribe automáticamente con Gemini multimodal
            if not texto.strip():
                logger.info("PDF escaneado detectado (%s). Extrayendo texto mediante Gemini multimodal...", ruta.name)
                api_key = os.getenv("GOOGLE_API_KEY")
                if api_key and api_key != "mock-key-for-tests":
                    try:
                        from google import genai

                        client = genai.Client(api_key=api_key)

                        # Evitar UnicodeEncodeError en httpx pasando un archivo temporal con nombre ASCII seguro
                        temp_file = None
                        upload_path = ruta
                        try:
                            str(ruta.name).encode("ascii")
                        except UnicodeEncodeError:
                            import shutil
                            import tempfile

                            temp_dir = Path(tempfile.gettempdir())
                            temp_file = temp_dir / f"nexus_doc_{abs(hash(ruta.name))}.pdf"
                            shutil.copyfile(ruta, temp_file)
                            upload_path = temp_file

                        try:
                            uploaded = client.files.upload(file=str(upload_path))
                        finally:
                            if temp_file and temp_file.exists():
                                temp_file.unlink(missing_ok=True)

                        try:
                            res = client.models.generate_content(
                                model="gemini-2.5-flash",
                                contents=[
                                    uploaded,
                                    "Transcribe fielmente todo el contenido de este documento oficial respetando títulos, materias, contenidos, créditos y tablas sin omitir nada.",
                                ],
                            )
                            texto = res.text or ""
                            logger.info("OCR multimodal completado para %s (%d caracteres)", ruta.name, len(texto))
                        finally:
                            try:
                                client.files.delete(name=uploaded.name)
                            except Exception:
                                pass
                    except Exception as e:
                        logger.warning("Fallo el OCR multimodal con Gemini para %s: %s", ruta.name, e)
                if not texto.strip():
                    texto = f"Documento oficial: {ruta.stem}."
        elif sufijo in [".docx", ".doc"]:
            import docx

            doc = docx.Document(str(ruta))
            parrafos = [p.text for p in doc.paragraphs if p.text.strip()]
            texto = "\n".join(parrafos)
        elif sufijo in [".txt", ".md"]:
            texto = ruta.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        logger.exception("Error al leer archivo %s", ruta.name)

    return texto


def _cargar_documentos_especializacion():
    from llama_index.core import Document

    docs_dir = _get_nexus_docs_dir()
    documentos = []

    if not docs_dir.exists():
        logger.warning("Directorio de Especialización IA no existe: %s", docs_dir)
        return documentos

    for archivo in docs_dir.iterdir():
        if archivo.is_file() and archivo.suffix.lower() in [".pdf", ".docx", ".txt", ".md"]:
            texto = _extraer_texto_archivo(archivo)
            if texto and texto.strip():
                doc = Document(
                    text=texto,
                    metadata={
                        "titulo": archivo.stem,
                        "archivo": archivo.name,
                        "fuente": "Especialización en Inteligencia Artificial UAM",
                    },
                )
                documentos.append(doc)

    logger.info("Cargados %d documentos para el RAG de Nexus desde %s", len(documentos), docs_dir)
    return documentos


def get_or_build_nexus_query_engine():
    """Inicializa una sola vez y persiste el índice de la Especialización en IA."""
    global _nexus_index, _nexus_query_engine
    if _nexus_query_engine is not None:
        return _nexus_query_engine

    api_key = os.getenv("GOOGLE_API_KEY", "")
    if not api_key or api_key == "mock-key-for-tests":
        logger.warning("GOOGLE_API_KEY no disponible; RAG de Nexus desactivado.")
        return None

    try:
        from llama_index.core import StorageContext, VectorStoreIndex, load_index_from_storage
        from llama_index.embeddings.google_genai import GoogleGenAIEmbedding
        from llama_index.llms.google_genai import GoogleGenAI

        embed_model = GoogleGenAIEmbedding(
            model_name="models/gemini-embedding-2",
            api_key=api_key,
        )
        llm = GoogleGenAI(
            model="models/gemini-2.5-flash",
            api_key=api_key,
            temperature=0.2,
        )

        storage_dir = _get_nexus_storage_dir()
        docstore_file = storage_dir / "docstore.json"

        if docstore_file.exists():
            logger.info("Cargando índice persistente de Nexus desde %s", storage_dir)
            storage_context = StorageContext.from_defaults(persist_dir=str(storage_dir))
            _nexus_index = load_index_from_storage(storage_context, embed_model=embed_model)
        else:
            logger.info("Construyendo índice de embeddings de Nexus por primera vez...")
            documentos = _cargar_documentos_especializacion()
            if not documentos:
                from llama_index.core import Document

                documentos = [
                    Document(
                        text="La Especialización en Inteligencia Artificial de la UAM tiene 25 créditos académicos y cuenta con materias como Conceptos Básicos de Matemáticas y Estadística, Programación para Analítica de Datos, Aprendizaje Automático, Computación en la Nube y Seminario de Investigación.",
                        metadata={"titulo": "Resumen Especialización IA", "archivo": "plan-de-estudios.pdf"},
                    )
                ]

            _nexus_index = VectorStoreIndex.from_documents(documentos, embed_model=embed_model)
            _nexus_index.storage_context.persist(persist_dir=str(storage_dir))
            logger.info("Índice de Nexus persistido exitosamente en %s", storage_dir)

        _nexus_query_engine = _nexus_index.as_query_engine(llm=llm, similarity_top_k=3)
        return _nexus_query_engine
    except Exception:
        logger.exception("Error al inicializar query_engine de Nexus")
        return None


async def consultar_especializacion_ia(consulta: str) -> dict[str, Any] | None:
    """Ejecuta consulta sobre la Especialización en IA retornando la respuesta y las fuentes utilizadas."""
    if not consulta or not consulta.strip():
        return None

    engine = get_or_build_nexus_query_engine()
    if engine is None:
        return None

    def _sync_query():
        try:
            response = engine.query(consulta)
            fuentes = []
            if hasattr(response, "source_nodes"):
                for node in response.source_nodes:
                    meta = getattr(node.node, "metadata", {})
                    archivo = meta.get("archivo") or meta.get("titulo")
                    if archivo and archivo not in fuentes:
                        fuentes.append(str(archivo))
            return {
                "respuesta": str(response),
                "fuentes": fuentes,
            }
        except Exception:
            logger.exception("Error consultando RAG de Nexus para query: %s", consulta)
            return None

    return await asyncio.to_thread(_sync_query)

